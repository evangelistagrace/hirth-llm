import os
import re
import textwrap
from pathlib import Path
from typing import Generator


import cohere
import chromadb
from chromadb.config import Settings
from dotenv import load_dotenv


load_dotenv()

COHERE_API_KEY = os.environ["COHERE_API_KEY"]
CHROMA_PATH = Path(__file__).parent / "chroma_db"
COLLECTION_NAME = "hirth_docs"
CHUNK_SIZE = 500   # tokens (approx chars / 4)
CHUNK_OVERLAP = 50

co = cohere.Client(COHERE_API_KEY)
chroma = chromadb.PersistentClient(path=str(CHROMA_PATH), settings=Settings(anonymized_telemetry=False))


def get_collection():
    return chroma.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )


# ── Parsers ────────────────────────────────────────────────────────────────────

def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def _parse_image(path: Path) -> str:
    import pytesseract
    from PIL import Image
    img = Image.open(str(path))
    # try DE+EN OCR
    return pytesseract.image_to_string(img, lang="deu+eng")


def _parse_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append(text)
    return "\n".join(lines)


def _parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    elif suffix in (".doc", ".docx"):
        return _parse_docx(path)
    elif suffix in (".xls", ".xlsx"):
        return _parse_xlsx(path)
    elif suffix in (".ppt", ".pptx"):
        return _parse_pptx(path)
    elif suffix in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
        return _parse_image(path)
    elif suffix in (".txt", ".md"):
        return path.read_text(errors="replace")
    else:
        return ""


# ── Chunking ───────────────────────────────────────────────────────────────────

def _split_chunks(text: str, source: str) -> Generator[dict, None, None]:
    """Sliding window over ~500-token chunks with 50-token overlap."""
    # rough token ≈ 4 chars
    size = CHUNK_SIZE * 4
    overlap = CHUNK_OVERLAP * 4
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return
    start = 0
    idx = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunk = text[start:end]
        yield {
            "text": chunk,
            "source": source,
            "chunk_index": idx,
        }
        idx += 1
        start += size - overlap


# ── Embedding ──────────────────────────────────────────────────────────────────

def _embed(texts: list[str], input_type: str = "search_document") -> list[list[float]]:
    response = co.embed(
        texts=texts,
        model="embed-multilingual-v3.0",
        input_type=input_type,
        embedding_types=["float"],
    )
    return response.embeddings.float


# ── Auto-categorisation ────────────────────────────────────────────────────────

CATEGORIES = ["mechanics", "electrics", "simulation", "software", "other"]

def _classify_document(filename: str, sample_text: str) -> str:
    """Ask GPT-4o-mini to classify the document into one of the fixed categories."""
    from openai import OpenAI
    client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])
    prompt = (
        f"Classify this engineering document into exactly one category.\n"
        f"Categories: mechanics, electrics, simulation, software\n"
        f"If none fit, reply: other\n\n"
        f"Filename: {filename}\n"
        f"Content sample:\n{sample_text[:800]}\n\n"
        f"Reply with only the category name in lowercase."
    )
    resp = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}],
        max_tokens=10,
        temperature=0,
    )
    raw = resp.choices[0].message.content.strip().lower()
    return raw if raw in CATEGORIES else "other"


# ── Ingestion ──────────────────────────────────────────────────────────────────

def ingest_file(path: Path) -> int:
    """Parse, chunk, embed and store a single file. Returns chunk count."""
    text = _parse_file(path)
    if not text.strip():
        return 0

    chunks = list(_split_chunks(text, source=path.name))
    if not chunks:
        return 0

    # classify once per document using a sample of the first chunk
    category = _classify_document(path.name, chunks[0]["text"])

    collection = get_collection()
    batch_size = 96  # Cohere max per request
    total = 0
    for i in range(0, len(chunks), batch_size):
        batch = chunks[i : i + batch_size]
        texts = [c["text"] for c in batch]
        embeddings = _embed(texts, input_type="search_document")
        ids = [f"{path.stem}_{c['chunk_index']}" for c in batch]
        metadatas = [
            {"source": c["source"], "chunk_index": c["chunk_index"], "category": category}
            for c in batch
        ]

        collection.upsert(
            ids=ids,
            documents=texts,
            embeddings=embeddings,
            metadatas=metadatas,
        )
        total += len(batch)

    return total


def ingest_directory(directory: Path) -> dict[str, int]:
    """Ingest all supported files in a directory."""
    results = {}
    supported = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".txt", ".md"}
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() in supported:
            count = ingest_file(path)
            results[path.name] = count
    return results


# ── Query ──────────────────────────────────────────────────────────────────────

def query_collection(query: str, n_results: int = 5, categories: list[str] | None = None) -> list[dict]:
    collection = get_collection()
    q_embedding = _embed([query], input_type="search_query")[0]
    where = None
    if categories:
        where = {"category": {"$in": categories}} if len(categories) > 1 else {"category": categories[0]}
    results = collection.query(
        query_embeddings=[q_embedding],
        n_results=n_results,
        include=["documents", "metadatas", "distances"],
        where=where,
    )
    chunks = []
    for doc, meta, dist in zip(
        results["documents"][0],
        results["metadatas"][0],
        results["distances"][0],
    ):
        chunks.append({
            "text": doc,
            "source": meta.get("source", ""),
            "chunk_index": meta.get("chunk_index", 0),
            "category": meta.get("category", "other"),
            "score": round(1 - dist, 4),
        })
    return chunks


def get_categories() -> dict[str, int]:
    """Return category → document count from metadata."""
    collection = get_collection()
    data = collection.get(include=["metadatas"])
    seen: dict[str, set] = {}
    for meta in data["metadatas"]:
        cat = meta.get("category", "other")
        src = meta.get("source", "")
        seen.setdefault(cat, set()).add(src)
    return {cat: len(srcs) for cat, srcs in sorted(seen.items())}
